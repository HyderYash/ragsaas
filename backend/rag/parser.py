from pathlib import Path
import csv
import json

import fitz
from docx import Document


def parse_pdf(file_path: Path) -> list[dict]:
    """
    Extract text page-by-page from a PDF.
    """

    pages = []

    document = fitz.open(file_path)

    try:
        for page_number, page in enumerate(document):
            text = page.get_text("text").strip()

            if text:
                pages.append(
                    {
                        "text": text,
                        "page": page_number + 1,
                    }
                )
    finally:
        document.close()

    return pages


def parse_docx(file_path: Path) -> list[dict]:
    """
    Extract text from a DOCX file.
    """

    document = Document(file_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ).strip()

    if not text:
        return []

    return [
        {
            "text": text,
            "page": None,
        }
    ]


def parse_txt(file_path: Path) -> list[dict]:
    """
    Extract text from TXT/MD files.
    """

    text = file_path.read_text(
        encoding="utf-8",
        errors="ignore",
    ).strip()

    if not text:
        return []

    return [
        {
            "text": text,
            "page": None,
        }
    ]


def parse_csv(file_path: Path) -> list[dict]:
    """
    Extract readable rows from CSV files.
    """

    rows = []

    with file_path.open(
        newline="",
        encoding="utf-8",
        errors="ignore",
    ) as handle:
        reader = csv.DictReader(handle)

        if reader.fieldnames:
            for index, row in enumerate(reader, start=1):
                row_text = ", ".join(
                    f"{key}: {value}"
                    for key, value in row.items()
                    if value
                )

                if row_text:
                    rows.append(f"Row {index}: {row_text}")
        else:
            handle.seek(0)

            for index, row in enumerate(csv.reader(handle), start=1):
                row_text = ", ".join(
                    value
                    for value in row
                    if value
                )

                if row_text:
                    rows.append(f"Row {index}: {row_text}")

    text = "\n".join(rows).strip()

    if not text:
        return []

    return [
        {
            "text": text,
            "page": None,
        }
    ]


def parse_json(file_path: Path) -> list[dict]:
    """
    Extract pretty-printed JSON content.
    """

    raw_text = file_path.read_text(
        encoding="utf-8",
        errors="ignore",
    )

    data = json.loads(raw_text)
    text = json.dumps(
        data,
        indent=2,
        ensure_ascii=False,
    )

    return [
        {
            "text": text,
            "page": None,
        }
    ]


def parse_pptx(file_path: Path) -> list[dict]:
    """
    Extract text from PowerPoint slides when python-pptx is installed.
    """

    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ValueError(
            "PPTX support requires python-pptx. Install requirements.txt again."
        ) from exc

    presentation = Presentation(file_path)
    pages = []

    for index, slide in enumerate(presentation.slides, start=1):
        parts = []

        for shape in slide.shapes:
            text = getattr(shape, "text", "")

            if text and text.strip():
                parts.append(text.strip())

        if parts:
            pages.append(
                {
                    "text": "\n".join(parts),
                    "page": index,
                }
            )

    return pages


def parse_xlsx(file_path: Path) -> list[dict]:
    """
    Extract cell values from Excel workbooks when openpyxl is installed.
    """

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            "XLSX support requires openpyxl. Install requirements.txt again."
        ) from exc

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    pages = []

    try:
        for sheet in workbook.worksheets:
            lines = []

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value)
                    for value in row
                    if value not in (None, "")
                ]

                if values:
                    lines.append(" | ".join(values))

            text = "\n".join(lines).strip()

            if text:
                pages.append(
                    {
                        "text": f"Sheet: {sheet.title}\n{text}",
                        "page": None,
                    }
                )
    finally:
        workbook.close()

    return pages


def parse_document(file_path: Path) -> list[dict]:
    """
    Automatically choose the parser based on extension.
    """

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return parse_pdf(file_path)

    if extension == ".docx":
        return parse_docx(file_path)

    if extension in {".txt", ".md"}:
        return parse_txt(file_path)

    if extension == ".csv":
        return parse_csv(file_path)

    if extension == ".json":
        return parse_json(file_path)

    if extension == ".pptx":
        return parse_pptx(file_path)

    if extension == ".xlsx":
        return parse_xlsx(file_path)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )
